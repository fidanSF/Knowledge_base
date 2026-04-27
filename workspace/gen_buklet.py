#!/usr/bin/env python3
"""ВольтМаркет буклет — A4 landscape, 2 slides (лицо + оборот + купон)"""

from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import qrcode
import io

# ── Brand colors ──────────────────────────────────────────────────
GRAY   = RGBColor(0x7F, 0x7F, 0x7E)   # фирменный серый
ORANGE = RGBColor(0xD7, 0x59, 0x45)   # фирменный оранжевый
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x2C, 0x2C, 0x2B)   # тёмный фон слайда 1
DARK2  = RGBColor(0x3F, 0x3F, 0x3E)   # чуть светлее
LGRAY  = RGBColor(0xF2, 0xF2, 0xF2)   # светлый для карточек слайда 2
MGRAY  = RGBColor(0x55, 0x55, 0x54)   # карточки слайда 1
MUTED  = RGBColor(0xBB, 0xBB, 0xBB)
DIM    = RGBColor(0x88, 0x88, 0x88)

# A4 landscape
W = Cm(29.7)
H = Cm(21.0)

prs = Presentation()
prs.slide_width  = int(W)
prs.slide_height = int(H)
BLANK = prs.slide_layouts[6]


# ── helpers ───────────────────────────────────────────────────────
def R(sl, x, y, w, h, fill, stroke=None, stroke_pt=1.5):
    """Добавляет прямоугольник на слайд."""
    s = sl.shapes.add_shape(1, int(x), int(y), int(w), int(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if stroke:
        s.line.color.rgb = stroke
        s.line.width = Pt(stroke_pt)
    else:
        s.line.fill.background()
    return s


def T(sl, text, x, y, w, h, size, bold=False, italic=False,
      color=WHITE, align=PP_ALIGN.LEFT, wrap=True):
    """Добавляет текстовый блок."""
    tb = sl.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.margin_left = Pt(3); tf.margin_right = Pt(3)
    tf.margin_top  = Pt(2); tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run()
    r.text = text; r.font.name = "Calibri"
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    return tb


def multi(sl, lines, x, y, w, h, def_align=PP_ALIGN.LEFT):
    """Несколько строк с разным форматированием."""
    tb = sl.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top  = Pt(2); tf.margin_bottom = Pt(2)
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ln.get('align', def_align)
        if 'sp' in ln:
            p.space_before = Pt(ln['sp'])
        r = p.add_run()
        r.text = ln.get('text', '')
        r.font.name = "Calibri"
        r.font.size = Pt(ln.get('sz', 11))
        r.font.bold = ln.get('b', False)
        r.font.italic = ln.get('i', False)
        r.font.color.rgb = ln.get('c', WHITE)
    return tb


# ── QR-код ────────────────────────────────────────────────────────
qr = qrcode.QRCode(version=1,
                   error_correction=qrcode.constants.ERROR_CORRECT_H,
                   box_size=10, border=2)
qr.add_data("https://volt-market.com")
qr.make(fit=True)
qr_img = qr.make_image(fill_color=(215, 89, 69), back_color=(255, 255, 255))
qr_buf = io.BytesIO()
qr_img.save(qr_buf, "PNG")


# ══════════════════════════════════════════════════════════════════
#  СЛАЙД 1 — ЛИЦЕВАЯ СТОРОНА  (Front)
# ══════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)

# Тёмный фон
R(s1, 0, 0, W, H, DARK)

# Оранжевые полосы сверху/снизу
BAR = Cm(1.15)
R(s1, 0, 0, W, BAR, ORANGE)
R(s1, 0, H - BAR, W, BAR, ORANGE)

# ── ЛЕВАЯ КОЛОННА (лого + слоган) ─────────────────────────────────
LW = Cm(10.2)   # ширина левой колонны

R(s1, 0, BAR, LW, H - 2 * BAR, GRAY)

# Логотип
T(s1, "ВОЛЬТ",  Cm(0.3), Cm(2.3), LW - Cm(0.5), Cm(2.0),
  44, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
T(s1, "МАРКЕТ", Cm(0.3), Cm(4.1), LW - Cm(0.5), Cm(1.5),
  30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# Разделитель
R(s1, Cm(1.2), Cm(5.9), LW - Cm(2.4), Cm(0.06), ORANGE)

# Слоган
T(s1, "делаем профессиональную\nэлектрику доступной",
  Cm(0.5), Cm(6.2), LW - Cm(1.0), Cm(2.2),
  11.5, italic=True, color=WHITE, align=PP_ALIGN.CENTER)

# Декор — три оранжевых полоски (символ молнии / энергии)
for i, (bw, bx) in enumerate([(Cm(1.6), Cm(1.0)),
                               (Cm(2.2), Cm(0.7)),
                               (Cm(1.4), Cm(1.3))]):
    R(s1, bx, Cm(8.9) + i * Cm(0.55), bw, Cm(0.35), ORANGE)

# Телефон и сайт
R(s1, Cm(0.6), H - Cm(3.6), LW - Cm(1.2), Cm(0.85), ORANGE)
T(s1, "8 800 550 11 61",
  Cm(0.6), H - Cm(3.6), LW - Cm(1.2), Cm(0.85),
  13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
T(s1, "volt-market.com",
  Cm(0.6), H - Cm(2.7), LW - Cm(1.2), Cm(0.6),
  10, color=MUTED, align=PP_ALIGN.CENTER)
T(s1, "звонок бесплатный",
  Cm(0.6), H - Cm(2.1), LW - Cm(1.2), Cm(0.5),
  8, italic=True, color=DIM, align=PP_ALIGN.CENTER)

# ── ПРАВАЯ СЕКЦИЯ — 5 карточек товаров ────────────────────────────
RX = LW
RW = W - LW

products = [
    ("РОЗЕТКИ И\nВЫКЛЮЧАТЕЛИ",   "Надёжные решения для вашего дома"),
    ("СВЕТИЛЬНИКИ\nНА ШИНОПРОВОД", "Современное трековое освещение"),
    ("СВЕТОДИОДНАЯ\nЛЕНТА",        "Яркий свет — низкое энергопотребление"),
    ("КАБЕЛЬ ВВГ\n(синий)",         "Силовые кабели высшего качества"),
    ("АВТОМАТЫ\nCHINT",            "Надёжная защита вашей электросети"),
]

COLS = 2
GAP_X = Cm(0.25)
GAP_Y = Cm(0.2)
PAD_X = Cm(0.35)  # отступ от края правой секции
PAD_Y = Cm(0.3)
card_w = (RW - PAD_X * 2 - GAP_X * (COLS - 1)) / COLS
card_h = (H - 2 * BAR - PAD_Y * 2 - GAP_Y * 2) / 3

for i, (name, desc) in enumerate(products):
    col = i % COLS
    row = i // COLS
    cx = RX + PAD_X + col * (card_w + GAP_X)
    cy = BAR + PAD_Y + row * (card_h + GAP_Y)

    # Карточка
    R(s1, cx, cy, card_w, card_h, MGRAY)
    # Оранжевый акцент слева
    R(s1, cx, cy, Cm(0.22), card_h, ORANGE)
    # Название товара
    T(s1, name, cx + Cm(0.35), cy + Cm(0.15), card_w - Cm(0.45), card_h * 0.48,
      9, bold=True, color=WHITE)
    # Описание
    T(s1, desc, cx + Cm(0.35), cy + Cm(0.15) + card_h * 0.48, card_w - Cm(0.45), card_h * 0.45,
      7.5, color=MUTED)

# Последняя (6-я) ячейка — CTA блок «Скидка до 15%»
col, row = 1, 2
cx = RX + PAD_X + col * (card_w + GAP_X)
cy = BAR + PAD_Y + row * (card_h + GAP_Y)

R(s1, cx, cy, card_w, card_h, ORANGE)
T(s1, "СКИДКА\nДО 15%",
  cx + Cm(0.3), cy + Cm(0.2), card_w - Cm(0.6), card_h * 0.58,
  16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
T(s1, "покажите этот буклет\nв любом магазине",
  cx + Cm(0.3), cy + card_h * 0.62, card_w - Cm(0.6), card_h * 0.32,
  8, italic=True, color=WHITE, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════
#  СЛАЙД 2 — ОБОРОТНАЯ СТОРОНА + КУПОН  (Back)
# ══════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)

# Белый фон
R(s2, 0, 0, W, H, WHITE)
R(s2, 0, 0, W, BAR, ORANGE)
R(s2, 0, H - BAR, W, BAR, GRAY)

# ── ЛЕВАЯ ЧАСТЬ — магазины + контакты (≈60% ширины) ──────────────
INFO_W = W * 0.60

# Заголовок
T(s2, "НАШИ МАГАЗИНЫ", Cm(0.7), Cm(1.5), Cm(13), Cm(1.1),
  20, bold=True, color=GRAY)

stores = [
    ("г. Оренбург", "пр. Победы, 151"),
    ("г. Оренбург", "пр. Автоматики, 28А"),
    ("г. Орск",     "пр. Ленина, 95А"),
]
each_w = (INFO_W - Cm(0.8)) / 3 - Cm(0.2)
for i, (city, addr) in enumerate(stores):
    sx = Cm(0.5) + i * (each_w + Cm(0.2))
    sy = Cm(3.0)
    R(s2, sx, sy, each_w, Cm(3.0), LGRAY)
    R(s2, sx, sy, Cm(0.22), Cm(3.0), ORANGE)
    T(s2, city, sx + Cm(0.4), sy + Cm(0.25), each_w - Cm(0.5), Cm(0.65),
      10.5, bold=True, color=ORANGE)
    T(s2, addr, sx + Cm(0.4), sy + Cm(0.95), each_w - Cm(0.5), Cm(0.9),
      9.5, color=GRAY, wrap=True)
    T(s2, "Пн–Вс: 9:00–20:00",
      sx + Cm(0.4), sy + Cm(1.9), each_w - Cm(0.5), Cm(0.55),
      8, italic=True, color=DIM)
    T(s2, "электротовары / освещение",
      sx + Cm(0.4), sy + Cm(2.5), each_w - Cm(0.5), Cm(0.45),
      7.5, color=DIM)

# Телефон
T(s2, "8 800 550 11 61",
  Cm(0.7), Cm(6.8), Cm(9), Cm(1.0),
  22, bold=True, color=ORANGE)
T(s2, "звонок бесплатный  |  volt-market.com",
  Cm(0.7), Cm(7.75), Cm(12), Cm(0.6),
  9, italic=True, color=DIM)

# Почему мы
T(s2, "ПОЧЕМУ ВЫБИРАЮТ НАС", Cm(0.7), Cm(8.7), Cm(14), Cm(0.85),
  14, bold=True, color=GRAY)

benefits = [
    "Широкий ассортимент электротоваров",
    "Профессиональные консультации бесплатно",
    "Цены от производителя — без наценок",
    "Только проверенные бренды и сертификаты",
    "Гарантия на все товары",
]
for i, b in enumerate(benefits):
    bx = Cm(0.7)
    by = Cm(9.8) + i * Cm(1.0)
    R(s2, bx, by + Cm(0.07), Cm(0.52), Cm(0.52), ORANGE)
    T(s2, "v",   # checkmark
      bx + Cm(0.06), by - Cm(0.02), Cm(0.52), Cm(0.62),
      10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    T(s2, b, bx + Cm(0.72), by, Cm(12), Cm(0.62), 10, color=GRAY)

# Статистика — вместо пустого места
STAT_Y = Cm(15.2)
stats = [
    ("11 000+", "наименований\nсветильников"),
    ("3 000+",  "моделей розеток\nи выключателей"),
    ("100%",    "оригинальные\nтовары"),
]
stat_w = (INFO_W - Cm(1.0)) / 3
for i, (num, lbl) in enumerate(stats):
    stx = Cm(0.5) + i * stat_w
    R(s2, stx + Cm(0.05), STAT_Y, stat_w - Cm(0.1), Cm(2.8), LGRAY)
    R(s2, stx + Cm(0.05), STAT_Y, stat_w - Cm(0.1), Cm(0.07), ORANGE)
    T(s2, num, stx, STAT_Y + Cm(0.2), stat_w, Cm(1.2),
      22, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
    T(s2, lbl, stx, STAT_Y + Cm(1.4), stat_w, Cm(1.2),
      8, italic=True, color=DIM, align=PP_ALIGN.CENTER)

# Слоган внизу
T(s2, "«Вместе создаём светлое будущее»",
  Cm(0.7), H - Cm(2.2), Cm(14), Cm(0.65),
  9.5, italic=True, color=DIM)

# ── ПРАВАЯ ЧАСТЬ — QR-код + купон ─────────────────────────────────
RQX = INFO_W + Cm(0.2)
RQW = W - RQX - Cm(0.3)

# QR-блок
T(s2, "Сканируйте — узнайте цены\nи ассортимент прямо сейчас",
  RQX, Cm(1.5), RQW, Cm(1.0), 9.5, color=GRAY, align=PP_ALIGN.CENTER)
qr_buf.seek(0)
QR_SZ = Cm(4.8)
s2.shapes.add_picture(qr_buf,
                      int(RQX + (RQW - QR_SZ) / 2),
                      int(Cm(2.7)),
                      int(QR_SZ), int(QR_SZ))
T(s2, "volt-market.com",
  RQX, Cm(7.7), RQW, Cm(0.7),
  11, bold=True, color=GRAY, align=PP_ALIGN.CENTER)

# ── КУПОН ─────────────────────────────────────────────────────────
CPX = RQX
CPY = Cm(8.7)
CPW = RQW
CPH = H - CPY - BAR - Cm(0.3)

# Рамка купона
R(s2, CPX, CPY, CPW, CPH, WHITE, stroke=ORANGE, stroke_pt=1.8)
# Шапка купона
R(s2, CPX, CPY, CPW, Cm(1.1), ORANGE)
T(s2, "КУПОН НА СКИДКУ",
  CPX + Cm(0.2), CPY + Cm(0.1), CPW - Cm(0.4), Cm(0.9),
  13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# Большая скидка
T(s2, "ДО 15%",
  CPX + Cm(0.2), CPY + Cm(1.15), CPW - Cm(0.4), Cm(1.8),
  40, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)

# Пояснение
multi(s2, [
    {'text': "скидка на любой товар", 'sz': 9.5, 'c': GRAY,
     'align': PP_ALIGN.CENTER},
    {'text': "при предъявлении этого буклета", 'sz': 9.5, 'c': GRAY,
     'align': PP_ALIGN.CENTER, 'sp': 1},
], CPX + Cm(0.2), CPY + Cm(3.0), CPW - Cm(0.4), Cm(1.0))

# Доп. бонус
R(s2, CPX + Cm(0.4), CPY + Cm(4.15), CPW - Cm(0.8), Cm(0.7),
  LGRAY)
T(s2, "Бесплатная консультация электрика",
  CPX + Cm(0.4), CPY + Cm(4.15), CPW - Cm(0.8), Cm(0.7),
  8.5, color=GRAY, align=PP_ALIGN.CENTER)

# Условия
T(s2, "* Не суммируется с другими акциями.\nДействителен до 31.12.2026",
  CPX + Cm(0.3), CPY + CPH - Cm(0.85), CPW - Cm(0.6), Cm(0.8),
  6.5, italic=True, color=DIM, align=PP_ALIGN.CENTER)


# ── Сохранение ────────────────────────────────────────────────────
OUT = r"D:\Desktop\ИИ\МЫСЛИ ПРОЕКТ ИИ\Буклет_ВольтМаркет.pptx"
prs.save(OUT)
print("OK:", OUT)
